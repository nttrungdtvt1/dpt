"""Build Chapters 2–3 Word report for the Per-Title Encoding thesis."""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
TOOLS = Path(__file__).resolve().parent
if str(TOOLS) not in sys.path:
    sys.path.insert(0, str(TOOLS))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from chapter2_body import write_chapter2
from chapter3_body import write_chapter3
from docx_report_utils import (
    add_table,
    caption,
    configure_styles,
    heading,
    note,
    page_break,
    para,
    set_run_font,
    setup_page,
    add_toc_field,
)

FIG = ROOT / "docs" / "chapter23_figures"
OUT_DOCS = ROOT / "docs"
OUT_ROOT = ROOT.parent  # D:\DPT
FILENAME = "Chuong_2_3_Co_so_ly_thuyet_va_Phuong_phap_Per_Title_Encoding.docx"


def cover(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TÀI LIỆU THÀNH PHẦN BÁO CÁO")
    set_run_font(r, size=14, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PER-TITLE ENCODING")
    set_run_font(r, size=18, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TỐI ƯU HÓA BITRATE LADDER THEO TỪNG NỘI DUNG VIDEO")
    set_run_font(r, size=16, bold=True)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CHƯƠNG 2. CƠ SỞ LÝ THUYẾT VÀ CÁC CHỈ SỐ ĐÁNH GIÁ")
    set_run_font(r, size=14, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CHƯƠNG 3. PHƯƠNG PHÁP NGHIÊN CỨU VÀ QUY TRÌNH ĐỀ XUẤT")
    set_run_font(r, size=14, bold=True)

    doc.add_paragraph()
    lines = [
        "Phạm vi: demo thực nghiệm nhỏ bằng FFmpeg và Python.",
        "Codec: H.264/AVC. Chỉ số chất lượng chính: VMAF.",
        "Không mô tả hệ thống VOD/OTT hoàn chỉnh. Không tái tạo quy trình Netflix.",
        "File này chỉ gồm Chương 2 và Chương 3.",
        "Các ví dụ số trong Chương 2 là minh họa, không phải kết quả đo.",
    ]
    for text in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=13, italic=True)


def toc_page(doc: Document) -> None:
    heading(doc, "MỤC LỤC", 1)
    para(
        doc,
        "Mở tài liệu bằng Microsoft Word, bấm chuột phải vào trường mục lục bên dưới và chọn Update Field "
        "(Update entire table) để hiện số trang. Các tiêu đề cấp 1–3 được đưa vào mục lục tự động.",
        indent=False,
    )
    add_toc_field(doc)
    para(
        doc,
        "Dàn ý chính: Chương 2 (tham số encoding, ladder, VMAF, SI/TI, đường cong bitrate–quality, "
        "Pareto và Convex Hull, tiêu chí và KPI). Chương 3 (bài toán, candidate, điều kiện thực nghiệm, "
        "chuẩn hóa, encode, SI/TI, đường cong, lọc Pareto/hull, quy tắc chọn ladder, baseline, so sánh, flowchart).",
        indent=False,
    )


def extras(doc: Document) -> None:
    heading(doc, "DANH SÁCH HÌNH VÀ BẢNG CẦN BỔ SUNG / RÀ SOÁT", 1)
    para(
        doc,
        "Các hình minh họa trong file đã được tạo từ số liệu giả định (Chương 2) hoặc từ logic quy trình (Chương 3). "
        "Khi có kết quả đo thật, không thay các hình minh họa Chương 2 bằng số liệu thí nghiệm; số liệu thật thuộc Chương 5.",
    )
    add_table(
        doc,
        ["Mã", "Nội dung", "Trạng thái"],
        [
            ["Hình 2.1", "Đường cong bitrate–quality (minh họa)", "Đã chèn PNG minh họa"],
            ["Hình 2.2", "Pareto dominance (minh họa)", "Đã chèn PNG minh họa"],
            ["Hình 2.3", "Pareto khác upper convex envelope (minh họa)", "Đã chèn PNG minh họa"],
            ["Hình 3.1", "Mô hình I/O", "Đã chèn PNG"],
            ["Hình 3.2", "Flowchart quy trình nghiên cứu", "Đã chèn PNG; có thể vẽ lại trong Word"],
            ["Bảng 3.2", "Thông số video thử nghiệm chính", "Cần điền file thật 10–30 giây"],
            ["min_vmaf", "Ngưỡng chất lượng cấu hình", "Cần điền nếu khác mặc định 0"],
            ["Baseline", "Nếu đổi YAML", "Điền cấu hình baseline thực tế"],
        ],
        "Bảng P.1. Hạng mục cần rà soát trước khi ghép vào báo cáo đầy đủ",
    )

    heading(doc, "DANH SÁCH TÀI LIỆU THAM KHẢO CẦN XÁC MINH", 1)
    para(
        doc,
        "Các mục dưới đây được dùng làm cơ sở định nghĩa hoặc công cụ. Cần đối chiếu đường dẫn/năm xuất bản "
        "trước khi đưa vào danh mục chính thức của báo cáo. Không bổ sung bài báo chưa đọc.",
        indent=False,
    )
    refs = [
        "[1] Aaron, A., Li, Z., Manohara, M., De Cock, J., Ronca, D. Per-Title Encode Optimization. Netflix TechBlog. "
        "https://netflixtechblog.com/per-title-encode-optimization-7e99442b62a2",
        "[2] Li, Z. et al. Toward A Practical Perceptual Video Quality Metric. Netflix TechBlog. "
        "https://netflixtechblog.com/toward-a-practical-perceptual-video-quality-metric-653f208b9652",
        "[3] Netflix. VMAF: The Journey Continues. https://netflixtechblog.com/vmaf-the-journey-continues-44b51ee9ed12",
        "[4] Netflix. VMAF GitHub repository. https://github.com/Netflix/vmaf",
        "[5] ITU-T. Recommendation P.910: Subjective video quality assessment methods for multimedia applications. "
        "https://www.itu.int/rec/T-REC-P.910",
        "[6] ITU-T. H.264: Advanced video coding for generic audiovisual services. https://www.itu.int/rec/T-REC-H.264",
        "[7] Apple. HLS authoring specification for Apple devices. "
        "https://developer.apple.com/documentation/http-live-streaming/hls-authoring-specification-for-apple-devices",
        "[8] FFmpeg Project. FFmpeg Documentation. https://ffmpeg.org/ffmpeg.html",
        "[9] FFmpeg Project. Encode/H.264. https://trac.ffmpeg.org/wiki/Encode/H.264",
        "[10] Wang, Z., Bovik, A. C., Sheikh, H. R., Simoncelli, E. P. (2004). Image quality assessment: from error "
        "visibility to structural similarity. IEEE TIP, 13(4), 600–612.",
        "[11] Robitza, W. CRF Guide (x264, x265 and libvpx). https://slhck.info/video/2017/02/24/crf-guide.html",
        "[12] Katsavounidis, I. Dynamic Optimizer — A Perceptual Video Encoding Optimization Framework. Netflix TechBlog "
        "(chỉ để phân biệt Per-Title / Per-Shot; không triển khai trong đề tài).",
    ]
    for item in refs:
        para(doc, item, indent=False)

    heading(doc, "GHI CHÚ PHẠM VI (KHÔNG THUỘC CHƯƠNG 1)", 1)
    para(
        doc,
        "Nội dung Chương 2–3 bám pipeline: candidate → VMAF → Pareto → upper convex hull → quy tắc chọn ladder, "
        "so với ladder cố định. SI/TI không chọn ladder chính. Per-Shot không triển khai. Không có số liệu VMAF "
        "thực nghiệm trong file này.",
    )


def build() -> list[Path]:
    from build_chapters_2_3_figures import main as make_figs

    make_figs()
    doc = Document()
    configure_styles(doc)
    setup_page(doc)
    cover(doc)
    page_break(doc)
    toc_page(doc)
    page_break(doc)
    write_chapter2(doc, FIG)
    page_break(doc)
    write_chapter3(doc, FIG)
    page_break(doc)
    extras(doc)

    OUT_DOCS.mkdir(parents=True, exist_ok=True)
    paths = [OUT_DOCS / FILENAME, OUT_ROOT / FILENAME]
    # If ROOT.parent is DPT when script lives in per-title-encoding/tools
    for path in paths:
        try:
            path.parent.mkdir(parents=True, exist_ok=True)
            doc.save(str(path))
        except OSError:
            continue
    return [p for p in paths if p.is_file()]


def self_check(path: Path) -> dict:
    from docx import Document as D

    d = D(str(path))
    texts = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    joined = "\n".join(texts)
    required = [
        "CHƯƠNG 2. CƠ SỞ LÝ THUYẾT VÀ CÁC CHỈ SỐ ĐÁNH GIÁ",
        "CHƯƠNG 3. PHƯƠNG PHÁP NGHIÊN CỨU VÀ QUY TRÌNH ĐỀ XUẤT",
        "2.1.",
        "2.2.",
        "2.3.",
        "2.4.",
        "2.5.",
        "2.6.",
        "2.7.",
        "2.8.",
        "3.1.",
        "3.2.",
        "3.3.",
        "3.4.",
        "3.5.",
        "3.6.",
        "3.7.",
        "3.8.",
        "3.9.",
        "3.10.",
        "3.11.",
        "3.12.",
        "Bitrate Saving",
        "ΔVMAF",
        "Pareto",
        "Convex",
        "VMAF",
        "Spatial Information",
        "Temporal Information",
        "Quy tắc lựa chọn được đề xuất trong phạm vi thực nghiệm",
    ]
    missing = [item for item in required if item not in joined]
    forbidden_hits = []
    for bad in ["tối ưu tuyệt đối", "tái tạo toàn bộ quy trình Netflix", "tập dữ liệu lớn"]:
        if bad in joined.lower() or bad in joined:
            forbidden_hits.append(bad)
    return {
        "path": str(path),
        "paragraphs": len(texts),
        "tables": len(d.tables),
        "missing": missing,
        "forbidden": forbidden_hits,
        "has_ch1": "Chương 1." in joined and "Tổng quan đề tài" in joined,
        "has_ch4": "Chương 4. Xây dựng demo" in joined,
        "has_ch5_results": "testsrc_720p.mp4: SI =" in joined,
    }


if __name__ == "__main__":
    saved = build()
    print("SAVED:")
    for p in saved:
        print(p)
        info = self_check(p)
        print(info)
