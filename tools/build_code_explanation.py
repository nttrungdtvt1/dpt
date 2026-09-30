"""Build a Word document that explains the ACTUAL source code."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
PLOTS = ROOT / "results" / "plots"
RESULTS = ROOT / "results"


def set_run_font(run, name: str = "Times New Roman", size: int = 13, code: bool = False) -> None:
    font = "Consolas" if code else name
    run.font.name = font
    run.font.size = Pt(11 if code else size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "Times New Roman" if not code else font)
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)
    if code:
        run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


def heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 14 if level == 2 else 13)


def para(doc: Document, text: str, *, indent: bool = True, bold: bool = False) -> None:
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Cm(1)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.4
    run = p.add_run(text)
    run.bold = bold
    set_run_font(run)


def note(doc: Document, kind: str, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(f"[{kind}] ")
    run.bold = True
    set_run_font(run, size=12)
    run2 = p.add_run(text)
    set_run_font(run2, size=12)


def code_block(doc: Document, text: str, caption_text: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, code=True)
    if caption_text:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(caption_text)
        r.italic = True
        set_run_font(r, size=11)


def caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    set_run_font(r, size=11)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], title: str) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        set_run_font(run, size=11)
    for r_i, row in enumerate(rows, start=1):
        for c_i, value in enumerate(row):
            cell = table.rows[r_i].cells[c_i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            set_run_font(run, size=10)
    caption(doc, title)


def picture(doc: Document, path: Path, title: str) -> None:
    if not path.is_file():
        note(doc, "Chưa có file", f"Không tìm thấy {path}.")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(14.5))
    caption(doc, title)


def fn_box(doc: Document, name: str, items: list[tuple[str, str]]) -> None:
    heading(doc, name, 3)
    rows = [[k, v] for k, v in items]
    add_table(doc, ["Mục", "Nội dung thực tế trong code"], rows, f"Bảng mô tả {name}")


def _load_payload() -> dict:
    path = RESULTS / "experiment_results.json"
    if not path.is_file():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def _load_rows() -> list[dict]:
    path = RESULTS / "experiment_results.csv"
    if not path.is_file():
        return []
    with path.open(encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def build() -> Path:
    DOCS.mkdir(parents=True, exist_ok=True)
    payload = _load_payload()
    rows = _load_rows()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.0)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Giải thích code thực tế — Per-Title Encoding  |  Trang ")
    set_run_font(r, size=10)
    add_page_number(footer)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("TÀI LIỆU GIẢI THÍCH CODE\nPer-Title Bitrate Ladder Selection cho VOD")
    r.bold = True
    set_run_font(r, size=18)

    para(
        doc,
        "Tài liệu này mô tả đúng code trong per-title-encoding sau khi nâng cấp pipeline "
        "candidate → VMAF → Pareto → upper convex hull. Không mô tả kiến trúc giả định.",
        indent=False,
    )
    note(
        doc,
        "Đã implement",
        "Candidate grid YAML; encode bằng ffmpeg_encoder.encode_rung (không viết lại); "
        "VMAF libvmaf model version=vmaf_v0.6.1; PSNR/SSIM phụ; Pareto; upper convex hull "
        "(monotone chain); hull-based ladder selection; fixed ladder; heuristic k×fixed; SI/TI.",
    )
    note(
        doc,
        "Không làm",
        "Không rewrite encoder/FFprobe. Không xóa heuristic cũ. Không gọi Pareto là convex hull. "
        "Không gọi kết quả là tối ưu toàn cục. Không giả lập VMAF.",
    )

    heading(doc, "1. Giới thiệu project", 1)
    para(
        doc,
        "Project là demo sinh viên: lựa chọn bitrate ladder theo từng nội dung ở mức thử nghiệm. "
        "Entry: python src/main.py. Phương pháp chính: Hull-based Per-Title Ladder Selection. "
        "Hai baseline: fixed ladder (Apple HLS H.264 rút gọn) và heuristic build_per_title_ladder(k).",
    )

    heading(doc, "2. Mục tiêu của project", 1)
    para(
        doc,
        "Trong code hiện tại: (1) FFprobe; (2) SI/TI mô tả phức tạp; (3) lưới candidate từ YAML; "
        "(4) encode H.264; (5) đo VMAF/PSNR/SSIM; (6) Pareto rồi upper hull; (7) chọn ladder thực dụng; "
        "(8) so sánh fixed + heuristic; (9) CSV/JSON/PNG; (10) pytest.",
    )

    heading(doc, "3. Kiến trúc tổng thể", 1)
    para(
        doc,
        "Dataclass: VideoMeta, ComplexityResult, VideoFeatures, LadderSpec, CandidateSpec, AppConfig, "
        "LadderRung, EncodeResult, QualityScores, ExperimentRow, RQPoint, FFmpegTools. Logic là hàm thuần.",
    )
    code_block(
        doc,
        "src/main.py\n"
        "    → config.load_config\n"
        "    → pipeline.run_pipeline\n"
        "         → extract_features (FFprobe + SI/TI)\n"
        "         → build_candidate_rungs\n"
        "         → encode_rung + measure_quality  (VMAF bắt buộc)\n"
        "         → pareto_frontier → upper_convex_hull\n"
        "         → select_ladder_from_hull\n"
        "         → build_fixed_ladder + build_per_title_ladder(k)   [baseline]\n"
        "         → write_csv / bitrate_vmaf.csv / write_json / plot_results",
        "Hình 3.1. Quan hệ module thực tế",
    )

    heading(doc, "4. Luồng xử lý dữ liệu", 1)
    code_block(
        doc,
        "Input Video\n"
        "     ↓\n"
        "FFprobe (probe_video)\n"
        "     ↓\n"
        "SI/TI (analyze_complexity)  → mô tả nội dung; k chỉ cho heuristic\n"
        "     ↓\n"
        "Candidate grid (build_candidate_rungs)  → không upscale\n"
        "     ↓\n"
        "FFmpeg encode_rung từng candidate\n"
        "     ↓\n"
        "VMAF + PSNR + SSIM (measure_quality)\n"
        "     ↓\n"
        "bitrate_vmaf.csv  (một dòng = một encode thật)\n"
        "     ↓\n"
        "remove_dominated_points → Pareto frontier\n"
        "     ↓\n"
        "upper_convex_hull  (monotone chain, X=bitrate Y=VMAF)\n"
        "     ↓\n"
        "select_ladder_from_hull  (Hull-based Per-Title Ladder Selection)\n"
        "     ↓\n"
        "Fixed ladder + heuristic k×fixed  (cùng metric)\n"
        "     ↓\n"
        "CSV / JSON / bitrate_vmaf.png + plot cũ",
        "Hình 4.1. Luồng đang chạy trong pipeline.py",
    )

    from draw_flowcharts import draw_all
    from insert_flowchart_chapter import add_flowchart_chapter

    drawn = draw_all()
    add_flowchart_chapter(
        doc,
        {"main": drawn[0], "candidate": drawn[1], "pareto": drawn[2], "ladder": drawn[3]},
    )

    heading(doc, "5. Cấu trúc thư mục", 1)
    add_table(
        doc,
        ["Đường dẫn", "Vai trò thực tế"],
        [
            ["src/main.py", "CLI argparse"],
            ["src/config.py", "YAML → AppConfig, CandidateSpec, ladder_selection"],
            ["src/pipeline.py", "Điều phối pipeline khoa học"],
            ["src/analysis/rq_points.py", "RQPoint, Pareto, upper convex hull"],
            ["src/analyzer/", "FFprobe + SI/TI + CRF probe (tắt)"],
            ["src/ladder/", "Fixed / heuristic / candidate / hull / validate; logic ở models.py"],
            ["src/encoder/ffmpeg_encoder.py", "Encode libx264 — giữ nguyên"],
            ["src/metrics/", "VMAF/PSNR/SSIM + CSV/JSON"],
            ["src/visualization/plots.py", "Plot cũ + bitrate_vmaf.png"],
            ["config/config.yaml", "fixed_ladder, candidates, quality.vmaf true"],
            ["tests/", "23 test cũ + test Pareto/hull/candidate/VMAF/selection"],
            ["results/", "CSV, JSON, encodes, PNG"],
        ],
        "Bảng 5.1. Thư mục source",
    )

    heading(doc, "6. Giải thích từng file", 1)
    files = [
        ("src/main.py", "CLI: check-env, make-sample, analyze, run. Analyze vẫn in SI/TI + fixed + heuristic; hull cần lệnh run."),
        ("src/config.py", "Thêm CandidateSpec và ladder_selection; load_config đọc candidates từ YAML."),
        ("src/pipeline.py", "Thứ tự khoa học; cache encode theo (height, target bitrate); bắt buộc quality.vmaf."),
        ("src/analysis/rq_points.py", "remove_dominated_points, pareto_frontier, upper_convex_hull (monotone chain)."),
        ("src/analyzer/video_probe.py", "FFprobe JSON — không sửa."),
        ("src/analyzer/complexity.py", "SI Sobel, TI frame-diff — không sửa; vai trò đổi thành mô tả."),
        ("src/ladder/models.py", "Thêm build_candidate_rungs và select_ladder_from_hull; giữ build_per_title_ladder."),
        ("src/ladder/fixed_ladder.py / per_title_ladder.py / ladder_validator.py", "Wrapper ủy quyền, được giữ."),
        ("src/encoder/ffmpeg_encoder.py", "Không rewrite. _base_args list, libx264 -b:v, scale, GOP, -an."),
        ("src/metrics/quality_metrics.py", "VMAF bật thì raise nếu không parse được; ffmpeg_has_libvmaf."),
        ("src/visualization/plots.py", "Giữ 3 PNG cũ; thêm bitrate_vmaf.png (candidate, Pareto, hull, selected)."),
        ("config/config.yaml", "quality.vmaf true, vmaf_model version=vmaf_v0.6.1, candidates, ladder_selection."),
    ]
    for path, text in files:
        heading(doc, path, 2)
        para(doc, text)

    heading(doc, "7. Giải thích từng hàm quan trọng", 1)
    fn_box(
        doc,
        "probe_video(path, tools)",
        [
            ("File", "src/analyzer/video_probe.py"),
            ("Mục đích", "FFprobe JSON metadata."),
            ("Lệnh", "ffprobe -v error -print_format json -show_format -show_streams"),
        ],
    )
    fn_box(
        doc,
        "analyze_complexity(meta, tools, cfg)",
        [
            ("File", "src/analyzer/complexity.py"),
            ("Vai trò mới", "Mô tả SI/TI để đối chiếu đường cong; k chỉ cho heuristic."),
            ("Index", "C = clip(0.55*(SI/si_ref)+0.45*(TI/ti_ref), 0, 1.5)"),
        ],
    )
    fn_box(
        doc,
        "build_candidate_rungs(meta, cfg)",
        [
            ("File", "src/ladder/models.py"),
            ("Nguồn", "config.yaml candidates: height × bitrates_kbps"),
            ("Ràng buộc", "Bỏ height > nguồn; bitrate > 0; không trùng bitrate trong cùng height."),
        ],
    )
    fn_box(
        doc,
        "encode_rung(...)",
        [
            ("File", "src/encoder/ffmpeg_encoder.py — KHÔNG VIẾT LẠI"),
            ("Command", "scale bicubic, libx264, -b:v, -maxrate, -bufsize 2×, GOP ~2s, -an"),
        ],
    )
    fn_box(
        doc,
        "measure_quality(...)",
        [
            ("File", "src/metrics/quality_metrics.py"),
            ("VMAF", "libvmaf=model={vmaf_model}:n_threads=4; parse 'VMAF score:'; lỗi nếu thiếu điểm"),
            ("Phụ", "PSNR average:, SSIM All:; scale về độ phân giải nguồn"),
        ],
    )
    fn_box(
        doc,
        "pareto_frontier / remove_dominated_points",
        [
            ("File", "src/analysis/rq_points.py"),
            ("Luật", "B dominate A nếu B.bitrate ≤ A.bitrate AND B.VMAF ≥ A.VMAF và ít nhất một bất đẳng thức nghiêm"),
            ("Không phải", "Không phải convex hull"),
        ],
    )
    fn_box(
        doc,
        "upper_convex_hull(points)",
        [
            ("File", "src/analysis/rq_points.py"),
            ("Thuật toán", "Monotone chain, đi trái→phải, giữ vòng thuận kim đồng hồ (upper envelope)"),
            ("X, Y", "bitrate_kbps, vmaf; trùng bitrate giữ VMAF cao hơn; bỏ điểm collinear giữa"),
        ],
    )
    fn_box(
        doc,
        "select_ladder_from_hull(points, meta, cfg)",
        [
            ("Tên đúng", "Hull-based Per-Title Ladder Selection, không phải global optimum"),
            ("Tiêu chí", "min_vmaf; một điểm mỗi height trên hull (VMAF max); monotonic theo TARGET bitrate; max_representations"),
            ("Vì sao target", "Nội dung dễ nén có actual bitrate sụp xuống gần nhau; ABR cần bitrate danh định tăng dần"),
        ],
    )
    fn_box(
        doc,
        "build_per_title_ladder(meta, cfg, k)",
        [
            ("Vai trò", "Heuristic baseline k × fixed. KHÔNG XÓA."),
            ("Không phải", "Không phải hull, không phải Pareto"),
        ],
    )

    heading(doc, "8. Biến và dữ liệu", 1)
    add_table(
        doc,
        ["Tên", "Ý nghĩa"],
        [
            ["RQPoint.bitrate_kbps", "Bitrate thực đo (trục X Pareto/hull)"],
            ["RQPoint.vmaf", "VMAF libvmaf (trục Y)"],
            ["target_bitrate_kbps", "Bitrate danh định -b:v / bậc ladder"],
            ["method", "candidate | hull | fixed | per_title"],
            ["k / adjusted_k", "Chỉ heuristic"],
            ["si, ti, index", "Mô tả phức tạp nội dung"],
        ],
        "Bảng 8.1. Biến then chốt",
    )

    heading(doc, "9. FFprobe", 1)
    code_block(
        doc,
        "ffprobe -v error -print_format json -show_format -show_streams <video_path>",
        "probe_video — không shell",
    )

    heading(doc, "10. FFmpeg Encoding", 1)
    code_block(
        doc,
        "ffmpeg -y -hide_banner -loglevel error\n"
        "  -i <source>\n"
        "  -vf scale={width}:{height}:flags=bicubic\n"
        "  -c:v libx264 -preset veryfast\n"
        "  -b:v {bitrate}k -maxrate {bitrate}k -bufsize {2*bitrate}k\n"
        "  -g {gop} -keyint_min {gop} -sc_threshold 0\n"
        "  -pix_fmt yuv420p -an\n"
        "  <output.mp4>",
        "Khớp _base_args trong ffmpeg_encoder.py",
    )

    heading(doc, "11. VMAF Evaluation", 1)
    para(
        doc,
        "quality.vmaf: true. Filter: [0:v]scale=W:H:flags=bicubic,setsar=1[d];[1:v]setsar=1[r];"
        "[d][r]libvmaf=model=version=vmaf_v0.6.1:n_threads=4. Parse regex VMAF score:. "
        "Không invent điểm. Integration test skip có điều kiện nếu FFmpeg không có libvmaf; không biến fail thành pass.",
    )

    heading(doc, "12. Candidate grid và Rate–Quality Curve", 1)
    para(
        doc,
        "Lưới YAML, không hard-code trong Python. Mỗi candidate một encode thật. Đường cong chính: "
        "X = actual bitrate, Y = VMAF. PSNR/SSIM vẫn ghi CSV và vẽ bitrate_ssim.png.",
    )
    picture(doc, PLOTS / "bitrate_vmaf.png", "Hình 12.1. bitrate_vmaf.png từ số liệu thật")

    heading(doc, "13. Pareto Frontier", 1)
    para(
        doc,
        "Hàm pareto_frontier gọi remove_dominated_points. Pipeline: candidate → remove dominated → Pareto. "
        "Không đồng nghĩa convex hull. Test: empty, one, duplicate, dominated, non-dominated, equal bitrate, equal VMAF.",
    )

    heading(doc, "14. Upper Convex Hull", 1)
    para(
        doc,
        "Hàm upper_convex_hull trên Pareto. Monotone chain upper envelope. Xử lý rỗng, 1–2 điểm, collinear, "
        "trùng bitrate, NaN/Inf (bỏ). Không gọi thư viện hull. Không gọi Pareto là hull.",
    )

    heading(doc, "15. Per-Title Bitrate Ladder", 1)
    para(
        doc,
        "Phương pháp chính: select_ladder_from_hull. Không mặc định mọi điểm hull vào ABR. "
        "Heuristic build_per_title_ladder(k) vẫn chạy method=per_title để so sánh.",
    )

    heading(doc, "16. Fixed Ladder", 1)
    para(
        doc,
        "config.yaml: 360p@365, 432p@730, 720p@3000, 1080p@6000. Nguồn 720p bỏ 1080p. Encode và đo cùng VMAF/PSNR/SSIM.",
    )

    heading(doc, "17. So sánh ba phương pháp", 1)
    para(
        doc,
        "Fixed = baseline độc lập nội dung. Heuristic k×fixed = baseline cũ. Hull-based = phương pháp đề xuất. "
        "Cùng encoder và cùng metric. Không có công thức Saving % bắt buộc; dung lượng đọc từ size_bytes.",
    )
    picture(doc, PLOTS / "file_size_by_rung.png", "Hình 17.1. Dung lượng bậc ladder")
    picture(doc, PLOTS / "ladder_total_size.png", "Hình 17.2. Tổng dung lượng")
    picture(doc, PLOTS / "bitrate_ssim.png", "Hình 17.3. SSIM phụ")

    heading(doc, "18. Visualization", 1)
    para(
        doc,
        "Giữ file_size_by_rung.png, bitrate_ssim.png, ladder_total_size.png. Thêm bitrate_vmaf.png: "
        "chấm candidate, vòng Pareto, đường hull, sao ladder đã chọn. Cột size không dump toàn bộ candidate.",
    )

    heading(doc, "19. Test", 1)
    add_table(
        doc,
        ["File", "Nội dung"],
        [
            ["test_env/probe/complexity/ladder/ffmpeg_cmd/results/pipeline", "23 test cũ, giữ nguyên"],
            ["test_pareto.py", "empty, one, duplicate, dominated, non-dominated, equal bitrate/VMAF"],
            ["test_convex_hull.py", "empty/one/two, collinear, duplicate, upper envelope, NaN"],
            ["test_candidates.py", "YAML hợp lệ, no upscale, nhiều bitrate, invalid/duplicate"],
            ["test_ladder_selection.py", "hull points, monotonic, max reps, resolution, target vs actual"],
            ["test_vmaf.py", "parse regex; integration skip nếu không libvmaf"],
        ],
        "Bảng 19.1. Test",
    )
    para(doc, "Lần regression sau implementation: 54 passed (23 cũ + 31 mới).")

    heading(doc, "20. Luồng chạy", 1)
    code_block(
        doc,
        "python src\\main.py check-env\n"
        "python src\\main.py make-sample --output data\\samples\\testsrc_720p.mp4 --seconds 4 --size 1280x720 --pattern testsrc\n"
        "python src\\main.py run data\\samples\\testsrc_720p.mp4 data\\samples\\smptebars_720p.mp4 data\\samples\\testsrc2_720p.mp4 --output results\n"
        "python -m pytest -q\n"
        "python tools\\build_report.py\n"
        "python tools\\build_code_explanation.py",
    )

    heading(doc, "21. Một lần chạy thực tế", 1)
    if payload.get("analyses"):
        table_rows = []
        for item in payload["analyses"]:
            hull = ", ".join(f"{r['name']}@{r['bitrate_kbps']}" for r in item.get("hull_ladder") or [])
            table_rows.append(
                [
                    item["video"],
                    f"{item['complexity_label']} / {item['k']}",
                    f"SI {float(item['si']):.1f} TI {float(item['ti']):.1f}",
                    f"{len(item.get('pareto') or [])}/{len(item.get('convex_hull') or [])}",
                    hull or "—",
                ]
            )
        add_table(
            doc,
            ["Video", "Nhãn / k", "SI/TI", "Pareto/Hull", "Ladder hull (target)"],
            table_rows,
            "Bảng 21.1. Số liệu từ experiment_results.json",
        )
    para(
        doc,
        f"CSV có {len(rows)} dòng; mọi dòng lần chạy này có VMAF khác rỗng. "
        "Không viết kết luận ngoài những gì file kết quả chứa.",
    )

    heading(doc, "22. File output", 1)
    add_table(
        doc,
        ["File", "Tạo bởi"],
        [
            ["results/experiment_results.csv", "write_csv — gồm vmaf"],
            ["results/bitrate_vmaf.csv", "pipeline — chỉ candidate"],
            ["results/experiment_results.json", "analyses: rq_points, pareto, convex_hull, hull_ladder, fixed, per_title"],
            ["results/encodes/.../*.mp4", "encode_rung"],
            ["results/plots/bitrate_vmaf.png + 3 plot cũ", "plot_results"],
        ],
        "Bảng 22.1. Output",
    )

    heading(doc, "23. Giới hạn", 1)
    para(doc, "Clip lavfi ngắn. One-pass ABR lệch target trên nội dung dễ nén. Một điểm VMAF cao nhất mỗi height trên hull có thể bỏ bậc trung gian. Không HLS, không per-shot, không BD-Rate.")

    heading(doc, "24. Đề xuất cải thiện", 1)
    para(doc, "Clip camera thật; two-pass; ràng buộc ABR giữ bậc thấp cho mạng yếu; BD-Rate; HLS.")

    heading(doc, "25. Tổng kết", 1)
    para(
        doc,
        "Code hiện tại là demo hull-based per-title trên VMAF thật, giữ heuristic và fixed để so sánh. "
        "FFmpeg/FFprobe/encoder không bị viết lại. Pareto và hull được tách rõ trong src/analysis/rq_points.py.",
    )

    heading(doc, "26. Câu hỏi bảo vệ (theo code thực tế)", 1)
    qa = [
        ("Per-Title Encoding trong code là gì?", "Phương pháp chính: chọn ladder từ upper hull của Pareto bitrate–VMAF. Heuristic k×fixed chỉ là baseline."),
        ("Fixed Ladder là gì?", "Bảng YAML Apple 365/730/3000/6000 kbps, không upscale."),
        ("VMAF làm gì?", "Metric chính; libvmaf model version=vmaf_v0.6.1; parse VMAF score từ stderr FFmpeg."),
        ("Pareto Frontier là gì?", "Tập điểm không bị dominate (min bitrate, max VMAF). Hàm pareto_frontier."),
        ("Upper Convex Hull là gì?", "Bao lồi phía trên (bitrate, VMAF), monotone chain, hàm upper_convex_hull."),
        ("Pareto khác hull thế nào?", "Pareto = không dominated. Hull = tập con nằm trên envelope lồi. Pipeline làm Pareto trước rồi hull."),
        ("Ladder được chọn thế nào?", "select_ladder_from_hull: không lấy hết điểm hull; một height một điểm VMAF cao nhất; target bitrate tăng dần; cắt max_representations."),
        ("SI/TI còn dùng không?", "Có, để mô tả nội dung và giải thích khác biệt đường cong; không gán bitrate cho phương pháp chính."),
        ("FFmpeg/FFprobe làm gì?", "Encode H.264, filter VMAF/PSNR/SSIM, probe metadata. Python điều phối."),
        ("Có tối ưu toàn cục không?", "Không. Disclaimer JSON và docstring: hull-based selection, mức thử nghiệm."),
    ]
    for i, (q, a) in enumerate(qa, start=1):
        heading(doc, f"Câu {i}. {q}", 3)
        para(doc, a)

    out = DOCS / "Per_Title_Encoding_Code_Explanation.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    print(build())
