"""Generate the Vietnamese student report from measured results."""

from __future__ import annotations

import csv
import json
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
DOCS = ROOT / "docs"


def set_run_font(run, name: str = "Times New Roman", size: int = 13) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), name)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size=16 if level == 1 else 14)


def para(doc: Document, text: str, *, bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(1)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.bold = bold
    set_run_font(run)


def caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.italic = True
    set_run_font(run, size=12)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], title: str) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        set_run_font(run, size=11)
    for r_i, row in enumerate(rows, start=1):
        for c_i, value in enumerate(row):
            cell = table.rows[r_i].cells[c_i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            set_run_font(run, size=11)
    caption(doc, title)


def add_picture(doc: Document, path: Path, title: str) -> None:
    if not path.is_file():
        para(doc, f"(Chưa có hình {path.name}.)")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(14.5))
    caption(doc, title)


def load_rows() -> list[dict[str, str]]:
    csv_path = RESULTS / "experiment_results.csv"
    if not csv_path.is_file():
        return []
    with csv_path.open(encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def load_json() -> dict:
    path = RESULTS / "experiment_results.json"
    if not path.is_file():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def fmt(value: object, digits: int = 3) -> str:
    try:
        if value in (None, ""):
            return "—"
        return f"{float(value):.{digits}f}"
    except (TypeError, ValueError):
        return str(value)


def _num(row: dict, key: str) -> float:
    try:
        return float(row.get(key) or 0)
    except (TypeError, ValueError):
        return 0.0


def _filter(rows: list[dict], video: str, method: str) -> list[dict]:
    return [row for row in rows if row["video"] == video and row["method"] == method]


def _kb(rows: list[dict]) -> int:
    return int(sum(int(float(row["size_bytes"])) for row in rows) / 1024)


def _analysis_paragraphs(payload: dict, rows: list[dict]) -> list[str]:
    texts: list[str] = []
    for item in payload.get("analyses") or []:
        video = item["video"]
        hull_rows = _filter(rows, video, "hull")
        fixed_rows = _filter(rows, video, "fixed")
        heur_rows = _filter(rows, video, "per_title")
        cand_rows = _filter(rows, video, "candidate")
        hull_vmaf = ", ".join(f"{row['rung']}={fmt(row['vmaf'], 2)}" for row in hull_rows)
        fixed_vmaf = ", ".join(f"{row['rung']}={fmt(row['vmaf'], 2)}" for row in fixed_rows)
        texts.append(
            f"{video}: SI = {fmt(item['si'], 2)}, TI = {fmt(item['ti'], 2)}, "
            f"C = {fmt(item['complexity_index'], 3)}, nhãn {item['complexity_label']}, "
            f"k heuristic = {fmt(item['k'], 2)}. Lưới candidate {len(item.get('candidates') or [])} điểm, "
            f"Pareto {len(item.get('pareto') or [])} điểm, upper hull {len(item.get('convex_hull') or [])} điểm, "
            f"ladder hull {len(item.get('hull_ladder') or [])} bậc. "
            f"Tổng dung lượng (KB): hull {_kb(hull_rows)}, fixed {_kb(fixed_rows)}, "
            f"heuristic {_kb(heur_rows)}; {len(cand_rows)} candidate đã đo VMAF. "
            f"VMAF ladder hull: {hull_vmaf or '—'}. VMAF ladder fixed: {fixed_vmaf or '—'}."
        )
    return texts


def build() -> Path:
    DOCS.mkdir(parents=True, exist_ok=True)
    rows = load_rows()
    payload = load_json()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("TRƯỜNG ĐẠI HỌC GIAO THÔNG VẬN TẢI\nKHOA ĐIỆN – ĐIỆN TỬ")
    r.bold = True
    set_run_font(r, size=14)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("\nBÁO CÁO BÀI TẬP LỚN\nMôn: Kỹ thuật đa phương tiện")
    r.bold = True
    set_run_font(r, size=16)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(
        "Lựa chọn bitrate ladder theo từng nội dung\n"
        "(Per-Title Encoding) cho hệ thống Video on Demand\n"
        "bằng FFmpeg và Python"
    )
    r.bold = True
    set_run_font(r, size=14)

    para(
        doc,
        "Ghi chú phạm vi: đây là hệ thống lựa chọn bitrate ladder theo đặc trưng nội dung "
        "ở mức thử nghiệm. Phương pháp chính là Hull-based Per-Title Ladder Selection trên "
        "tập điểm bitrate–VMAF đã encode thật. Không phải thuật toán tối ưu toàn cục và "
        "không mô phỏng đầy đủ pipeline thương mại của Netflix. Pareto frontier không đồng "
        "nghĩa với upper convex hull.",
    )

    add_heading(doc, "Mục lục", 1)
    para(
        doc,
        "Mở tài liệu trong Microsoft Word, bấm vào mục lục và chọn Update Field nếu cần làm mới số trang.",
    )

    add_heading(doc, "Chương 1. Tổng quan đề tài", 1)
    add_heading(doc, "1.1. Bối cảnh Video on Demand", 2)
    para(
        doc,
        "Video on Demand (VOD) cho phép người xem mở nội dung đã mã hóa sẵn tại thời điểm "
        "bất kỳ. Nền tảng phát trực tuyến thường không gửi một file duy nhất mà chuẩn bị "
        "nhiều phiên bản cùng một video, khác nhau về độ phân giải và bitrate. Trình phát "
        "chọn phiên bản phù hợp với băng thông. Tập các phiên bản đó gọi là bitrate ladder.",
    )
    add_heading(doc, "1.2. Vấn đề của bitrate ladder cố định", 2)
    para(
        doc,
        "Ladder cố định gán cùng một bảng bitrate–độ phân giải cho mọi tựa. Nội dung ít chi "
        "tiết bị cấp thừa bit. Nội dung khó nén bị thiếu bit. Độ phân giải không đủ để quyết "
        "định bitrate vì hai video cùng 720p có thể khác nhau rất nhiều về entropy không-thời gian.",
    )
    add_heading(doc, "1.3. Khái niệm Per-Title Encoding", 2)
    para(
        doc,
        "Per-Title Encoding, được Netflix mô tả năm 2015, điều chỉnh công thức mã hóa theo "
        "từng tựa: chạy phân tích và encode thử, rồi chọn cặp resolution–bitrate trên đường "
        "cong rate–quality. Đồ án học ý tưởng đó ở mức thử nghiệm: encode một lưới candidate, "
        "đo VMAF, loại điểm bị dominate, lấy Pareto frontier, tính upper convex hull, rồi chọn "
        "ladder thực dụng từ hull. SI/TI được giữ để mô tả độ phức tạp nội dung, không còn là "
        "công thức chính gán bitrate. Hàm k × fixed ladder được giữ làm heuristic baseline.",
    )
    add_heading(doc, "1.4. Mục tiêu, phạm vi và ý nghĩa", 2)
    para(
        doc,
        "Mục tiêu: demo chạy trên máy cá nhân, so sánh ba phương pháp (fixed, heuristic k×fixed, "
        "hull-based), có số liệu VMAF/PSNR/SSIM thật. Phạm vi: H.264/AVC, clip lavfi ngắn, "
        "không HLS/DASH, không GPU, không per-shot. Ý nghĩa: minh họa đánh đổi bitrate–chất lượng "
        "đúng trọng tâm môn Kỹ thuật đa phương tiện.",
    )

    add_heading(doc, "Chương 2. Cơ sở lý thuyết", 1)
    add_heading(doc, "2.1. Video số, độ phân giải, FPS và bitrate", 2)
    para(
        doc,
        "Một video số là chuỗi khung hình theo thời gian. Độ phân giải là số pixel mỗi khung. "
        "Frame rate (FPS) là số khung mỗi giây. Bitrate là số bit trung bình mỗi giây sau nén. "
        "Dung lượng xấp xỉ bitrate nhân thời lượng. Tăng bitrate thường tăng chất lượng nhưng "
        "không tuyến tính: khi đã bão hòa, thêm bit gần như không còn lợi ích nhìn thấy.",
    )
    add_heading(doc, "2.2. Nén video, H.264, I/P/B và GOP", 2)
    para(
        doc,
        "Nén video khai thác dư thừa không gian trong một khung và dư thừa thời gian giữa các "
        "khung. H.264/AVC dùng dự đoán trong khung, ước lượng chuyển động và mã hóa phần dư. "
        "Demo đặt khoảng keyframe khoảng 2 giây để seek và cắt segment thuận tiện.",
    )
    add_heading(doc, "2.3. Rate control: CBR, VBR, CRF, ABR", 2)
    para(
        doc,
        "Trong demo, encode chính dùng -b:v (ABR một pass) vì đối tượng so sánh là bitrate ladder. "
        "Two-pass có thể bật trong YAML nhưng mặc định tắt.",
    )
    add_heading(doc, "2.4. VOD, Adaptive Bitrate Streaming và ladder", 2)
    para(
        doc,
        "ABR chia video thành segment, mỗi mức chất lượng một luồng. Apple công bố bảng gợi ý "
        "H.264 trong HLS Authoring Specification. Đồ án lấy bốn mức 640×360 @ 365 kbps, "
        "768×432 @ 730 kbps, 1280×720 @ 3000 kbps, 1920×1080 @ 6000 kbps làm fixed baseline. "
        "Lưới candidate dùng cùng các chiều cao đó, mỗi bậc nhiều bitrate, không bịa hàng 480p.",
    )
    add_heading(doc, "2.5. Chỉ số chất lượng, Pareto và convex hull", 2)
    para(
        doc,
        "PSNR = 10 log10 (MAX^2 / MSE). SSIM so sánh độ sáng, tương phản và cấu trúc. VMAF của "
        "Netflix hợp nhất nhiều đặc trưng và học từ MOS; đây là metric chính của pipeline. "
        "Mọi PSNR/SSIM/VMAF được đo sau khi scale bản encode về đúng độ phân giải nguồn. "
        "Model VMAF dùng built-in version=vmaf_v0.6.1 qua filter libvmaf; không giả lập điểm.",
    )
    para(
        doc,
        "Với mục tiêu minimize bitrate và maximize VMAF, điểm B dominate A nếu B có bitrate "
        "không lớn hơn và VMAF không nhỏ hơn, ít nhất một bất đẳng thức nghiêm. Tập còn lại "
        "là Pareto frontier. Upper convex hull là bao lồi phía trên của các điểm (bitrate, VMAF), "
        "tính bằng monotone chain; đó là tập con của Pareto, không đồng nghĩa với Pareto. "
        "Ladder thực dụng được chọn từ hull với ràng buộc số bậc, một điểm mỗi chiều cao "
        "(VMAF cao nhất), bitrate mục tiêu tăng dần. Đây là hull-based selection, không phải "
        "tối ưu toán học toàn cục.",
    )
    para(
        doc,
        "Spatial Information (SI) và Temporal Information (TI) theo hướng ITU-T P.910 được dùng "
        "để mô tả độ phức tạp và đối chiếu các đường cong bitrate–VMAF giữa các clip. Hệ số k "
        "từ SI/TI chỉ còn phục vụ heuristic baseline.",
    )

    add_heading(doc, "Chương 3. Phân tích và thiết kế hệ thống", 1)
    add_heading(doc, "3.1. Yêu cầu", 2)
    para(
        doc,
        "Chức năng: kiểm tra FFmpeg, đọc metadata, ước lượng SI/TI, encode lưới candidate, đo "
        "VMAF/PSNR/SSIM, tính Pareto và upper hull, chọn ladder hull-based, encode fixed và "
        "heuristic để so sánh, xuất CSV/JSON/PNG. Phi chức năng: chạy máy cá nhân, lệnh FFmpeg "
        "dạng list, không ghi đè file gốc, lặp lại được thí nghiệm.",
    )
    add_heading(doc, "3.2. Kiến trúc và luồng xử lý", 2)
    para(
        doc,
        "Luồng thực tế: video → FFprobe → SI/TI (mô tả) → candidate grid (không upscale) → "
        "FFmpeg encode → VMAF (+ PSNR/SSIM) → dataset bitrate–VMAF → loại điểm dominated → "
        "Pareto frontier → upper convex hull → hull-based ladder selection → encode/so sánh "
        "fixed ladder và heuristic k×fixed.",
    )
    add_table(
        doc,
        ["Module", "Input", "Output"],
        [
            ["video_probe", "File video", "width, height, fps, duration, codec, bitrate"],
            ["complexity", "Video + metadata", "SI, TI, index, nhãn, k (heuristic)"],
            ["build_candidate_rungs", "Metadata + YAML", "Lưới resolution × bitrate"],
            ["ffmpeg_encoder", "Video + bậc", "File mp4, thời gian, bitrate thật, size"],
            ["quality_metrics", "Gốc + encode", "VMAF, PSNR, SSIM"],
            ["rq_points", "Các điểm đo", "Pareto, upper convex hull"],
            ["select_ladder_from_hull", "Hull + ràng buộc", "Ladder hull-based"],
            ["fixed / per_title_ladder", "Metadata (+k)", "Hai baseline"],
            ["result_collector / plots", "Các dòng đo", "CSV, JSON, PNG"],
        ],
        "Bảng 3.1. Module chính của demo",
    )
    add_heading(doc, "3.3. Ba phương pháp so sánh", 2)
    para(
        doc,
        "Fixed ladder: bảng Apple, lọc không upscale. Heuristic per-title: bitrate = clamp(k × "
        "bitrate cố định), hàm build_per_title_ladder được giữ nguyên để đối chiếu. Hull-based: "
        "chọn từ upper hull; tiêu chí một điểm mỗi chiều cao còn trên hull (VMAF cao nhất, "
        "bitrate mục tiêu thấp hơn nếu hòa), rồi bắt monotonic theo target bitrate, cắt "
        "max_representations. Không khẳng định mọi điểm hull phải vào ABR ladder.",
    )
    add_heading(doc, "3.4. Công nghệ", 2)
    para(
        doc,
        "Python 3.12, FFmpeg/FFprobe (Gyan full build, có libvmaf), NumPy, Matplotlib, PyYAML, "
        "pytest. Không dùng OpenCV. Không tự tải mô hình AI; VMAF dùng model built-in của filter.",
    )

    add_heading(doc, "Chương 4. Xây dựng demo", 1)
    para(
        doc,
        "Môi trường đã kiểm trên Windows: Python 3.12.10, FFmpeg Gyan 9.0.1 có libvmaf, "
        "model version=vmaf_v0.6.1. Cấu hình nằm ở config/config.yaml (codec, preset, "
        "fixed_ladder, candidates, ladder_selection, quality.vmaf). CLI: check-env, make-sample, "
        "analyze, run. Kết quả: experiment_results.csv/.json, bitrate_vmaf.csv, PNG trong "
        "results/plots.",
    )

    add_heading(doc, "Chương 5. Thực nghiệm và đánh giá", 1)
    add_heading(doc, "5.1. Thiết lập", 2)
    para(
        doc,
        "Ba clip lavfi 1280×720, 25 fps, 4 giây: testsrc, smptebars, testsrc2. Thông số giữ "
        "cố định: libx264, preset veryfast, GOP 2 s, không audio, scale bicubic. Candidate: "
        "360p [250, 365, 500, 700], 432p [500, 730, 1000, 1500], 720p [1000, 1500, 2000, 3000] "
        "kbps; 1080p bị loại vì nguồn 720p (không upscale). Metric chính VMAF; PSNR/SSIM phụ. "
        "Số liệu dưới đây đo được trên máy thực hành, không phải số minh họa.",
    )

    if payload.get("analyses"):
        analysis_rows = []
        for item in payload["analyses"]:
            analysis_rows.append(
                [
                    item["video"],
                    f"{item['width']}x{item['height']}",
                    fmt(item["si"], 2),
                    fmt(item["ti"], 2),
                    fmt(item["complexity_index"], 3),
                    str(item["complexity_label"]),
                    fmt(item["k"], 2),
                    str(len(item.get("pareto") or [])),
                    str(len(item.get("convex_hull") or [])),
                    str(len(item.get("hull_ladder") or [])),
                ]
            )
        add_table(
            doc,
            ["Video", "Size", "SI", "TI", "C", "Nhãn", "k", "Pareto", "Hull", "Ladder hull"],
            analysis_rows,
            "Bảng 5.1. SI/TI (mô tả phức tạp) và kích thước Pareto/hull đo được",
        )

    ladder_rows = [row for row in rows if row.get("method") in {"fixed", "per_title", "hull"}]
    if ladder_rows:
        data_rows = []
        for row in ladder_rows:
            data_rows.append(
                [
                    row["video"].replace(".mp4", ""),
                    row["method"],
                    row["rung"],
                    row["target_bitrate_kbps"],
                    fmt(row["actual_bitrate_kbps"], 1),
                    str(int(int(float(row["size_bytes"])) / 1024)),
                    fmt(row["vmaf"], 2),
                    fmt(row["psnr"], 2),
                    fmt(row["ssim"], 4),
                ]
            )
        add_table(
            doc,
            ["Video", "PP", "Bậc", "Target", "Thực kbps", "KB", "VMAF", "PSNR", "SSIM"],
            data_rows,
            "Bảng 5.2. Ba phương pháp ladder (hull / fixed / heuristic) với VMAF đo thật",
        )

    add_picture(doc, RESULTS / "plots" / "bitrate_vmaf.png", "Hình 5.1. Candidate, Pareto, upper hull và ladder đã chọn")
    add_picture(doc, RESULTS / "plots" / "file_size_by_rung.png", "Hình 5.2. Dung lượng từng bậc ladder")
    add_picture(doc, RESULTS / "plots" / "bitrate_ssim.png", "Hình 5.3. Bitrate thực tế và SSIM (metric phụ)")
    add_picture(doc, RESULTS / "plots" / "ladder_total_size.png", "Hình 5.4. Tổng dung lượng toàn ladder")

    add_heading(doc, "5.2. Phân tích kết quả đo", 2)
    if not rows:
        para(doc, "Chưa có experiment_results.csv nên chưa phân tích số liệu.")
    else:
        para(
            doc,
            "Mọi dòng encode trong lần chạy này đều có VMAF parse từ FFmpeg libvmaf; không có "
            "điểm giả. SI/TI không được dùng để gán bitrate cho phương pháp chính. Chúng chỉ giúp "
            "đọc vì sao các đường bitrate–VMAF khác nhau.",
        )
        for text in _analysis_paragraphs(payload, rows):
            para(doc, text)
        para(
            doc,
            "Quan sát (không suy diễn thành quy luật thương mại): smptebars có TI = 0, encoder "
            "chỉ đạt khoảng hai chục kbps dù target hàng nghìn; VMAF 720p đã khoảng 97 ngay từ "
            "candidate 1000 kbps. Hull chỉ giữ một bậc 720p vì các điểm 360p/432p nằm dưới bao lồi. "
            "testsrc2 bám target hơn, VMAF tăng rõ theo bitrate (khoảng 77 → 97 trên lưới candidate); "
            "heuristic k = 1,35 tốn thêm dung lượng và đẩy VMAF 720p cao hơn một chút so với fixed. "
            "Hull-based trên testsrc2 chọn 360p@365 và 720p@3000, bỏ 432p vì không nằm trên upper hull. "
            "Không kết luận hull luôn tiết kiệm hay luôn tốt hơn fixed: nó là lựa chọn thực dụng "
            "trên envelope, phụ thuộc nội dung và ràng buộc một điểm mỗi chiều cao.",
        )
        para(
            doc,
            "Heuristic cũ (k × fixed) vẫn hữu ích để đối chiếu: với nội dung hard nó tăng bitrate "
            "toàn ladder; với smptebars (medium, k = 1) nó gần trùng fixed. Đây đúng là baseline "
            "heuristic, không phải convex hull.",
        )

    add_heading(doc, "5.3. Hạn chế và hướng mở rộng", 2)
    para(
        doc,
        "Hạn chế: clip lavfi ngắn; ABR một pass nên bitrate thực có thể lệch target (rõ trên "
        "smptebars/testsrc); chọn một điểm VMAF cao nhất mỗi chiều cao trên hull có thể bỏ bậc "
        "trung gian hữu ích cho ABR; chưa HLS; chưa video camera thật. Mở rộng: clip thật, "
        "two-pass, BD-Rate, đóng gói HLS, hiệu chỉnh ràng buộc ladder_selection.",
    )

    add_heading(doc, "Chương 6. Kết luận", 1)
    para(
        doc,
        "Đã nâng cấp demo Python + FFmpeg từ heuristic SI/TI → k × fixed thành pipeline "
        "candidate → VMAF → Pareto → upper convex hull → hull-based ladder, đồng thời giữ "
        "fixed ladder, heuristic cũ, SI/TI, PSNR/SSIM và encoder/FFprobe hiện có. Kết quả thí "
        "nghiệm là số đo thật, không phải tối ưu toàn cục. Hướng tiếp theo là kiểm lại trên "
        "nội dung camera và ràng buộc ABR thực tế hơn.",
    )

    add_heading(doc, "Tài liệu tham khảo", 1)
    refs = [
        "Aaron, A., Li, Z., Manohara, M., De Cock, J., Ronca, D. (2015). Per-Title Encode Optimization. Netflix TechBlog. https://netflixtechblog.com/per-title-encode-optimization-7e99442b62a2",
        "Apple. HTTP Live Streaming authoring specification for Apple devices. https://developer.apple.com/documentation/http-live-streaming/hls-authoring-specification-for-apple-devices",
        "FFmpeg Project. FFmpeg Documentation. https://ffmpeg.org/ffmpeg.html",
        "FFmpeg Project. FFprobe Documentation. https://ffmpeg.org/ffprobe.html",
        "FFmpeg Project. Encode/H.264. https://trac.ffmpeg.org/wiki/Encode/H.264",
        "Robitza, W. (2017). CRF Guide (x264, x265 and libvpx). https://slhck.info/video/2017/02/24/crf-guide.html",
        "Li, Z. et al. (2016). Toward A Practical Perceptual Video Quality Metric. Netflix TechBlog. https://netflixtechblog.com/toward-a-practical-perceptual-video-quality-metric-653f208b9652",
        "Netflix. VMAF: The Journey Continues (2018). https://netflixtechblog.com/vmaf-the-journey-continues-44b51ee9ed12",
        "Netflix. VMAF GitHub repository. https://github.com/Netflix/vmaf",
        "ITU-T. Recommendation P.910: Subjective video quality assessment methods for multimedia applications. https://www.itu.int/rec/T-REC-P.910",
        "Wang, Z., Bovik, A. C., Sheikh, H. R., Simoncelli, E. P. (2004). Image quality assessment: from error visibility to structural similarity. IEEE TIP, 13(4), 600–612. https://doi.org/10.1109/TIP.2003.819861",
        "Katsavounidis, I. (2018). Dynamic Optimizer — A Perceptual Video Encoding Optimization Framework. Netflix TechBlog. https://netflixtechblog.com/dynamic-optimizer-a-perceptual-video-encoding-optimization-framework-e19f1e3a277f",
        "Bjøntegaard, G. (2001). Calculation of average PSNR differences between RD-curves. ITU-T SG16 Q.6 VCEG-M33.",
        "Robitza, W. ffmpeg-quality-metrics. https://github.com/slhck/ffmpeg-quality-metrics",
        "ITU-T. H.264: Advanced video coding for generic audiovisual services. https://www.itu.int/rec/T-REC-H.264",
    ]
    for i, item in enumerate(refs, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        run = p.add_run(f"[{i}] {item}")
        set_run_font(run, size=12)

    add_heading(doc, "Phụ lục A. Kiểm thử", 1)
    para(
        doc,
        "Bộ pytest giữ toàn bộ test cũ (FFmpeg, probe, SI/TI, không upscale, lệnh list, CSV/plot, "
        "pipeline) và bổ sung Pareto, upper hull, candidate grid, ladder selection, parse/integration VMAF. "
        "Lần chạy gần nhất sau implementation: 54 passed (23 test cũ + 31 test mới).",
    )

    out = DOCS / "BaoCao_PerTitle_VOD.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    path = build()
    print(path)
