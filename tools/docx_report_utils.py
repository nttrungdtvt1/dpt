"""Shared formatting helpers for Vietnamese technical reports."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, Emu
from docx.enum.style import WD_STYLE_TYPE


def set_run_font(run, name: str = "Times New Roman", size: int = 13, bold: bool | None = None, italic: bool | None = None):
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)


def _set_style_font(style, size: int, bold: bool = False, center: bool = False, space_before: int = 12, space_after: int = 6):
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(size)
    font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = style.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.5
    pf.keep_with_next = True
    pf.keep_together = True
    if center:
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(13)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "Times New Roman")
    npf = normal.paragraph_format
    npf.line_spacing = 1.5
    npf.space_after = Pt(8)
    npf.widow_control = True

    _set_style_font(doc.styles["Heading 1"], 16, bold=True, center=True, space_before=18, space_after=12)
    _set_style_font(doc.styles["Heading 2"], 14, bold=True, space_before=14, space_after=8)
    _set_style_font(doc.styles["Heading 3"], 13, bold=True, space_before=12, space_after=6)
    _set_style_font(doc.styles["Heading 4"], 13, bold=True, space_before=10, space_after=6)
    try:
        toc_style = doc.styles["TOC Heading"]
        _set_style_font(toc_style, 16, bold=True, center=True)
    except KeyError:
        pass


def setup_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = hp.add_run("Per-Title Encoding – Chương 2 và Chương 3")
    set_run_font(run, size=10, italic=True)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("Trang ")
    set_run_font(run, size=11)
    add_page_number(fp)


def add_page_number(paragraph) -> None:
    run1 = paragraph.add_run()
    set_run_font(run1, size=11)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    run1._r.append(fld1)

    run2 = paragraph.add_run()
    set_run_font(run2, size=11)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    run2._r.append(instr)

    run3 = paragraph.add_run()
    set_run_font(run3, size=11)
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run3._r.append(fld2)


def add_toc_field(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_begin)

    run2 = p.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    run2._r.append(instr)

    run3 = p.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run3._r.append(fld_sep)

    run4 = p.add_run("Cập nhật mục lục trong Microsoft Word: chuột phải vào mục lục → Update Field.")
    set_run_font(run4, size=12, italic=True)

    run5 = p.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run5._r.append(fld_end)


def heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        size = 16 if level == 1 else 14 if level == 2 else 13
        set_run_font(run, size=size, bold=True)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.bold = True


def para(doc: Document, text: str, *, indent: bool = True, bold: bool = False, italic: bool = False, size: int = 13) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.first_line_indent = Cm(1.0) if indent else Cm(0)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.widow_control = True
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)


def formula(doc: Document, text: str, number: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_together = True
    run = p.add_run(text)
    set_run_font(run, size=13, italic=True)
    if number:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2.paragraph_format.space_after = Pt(8)
        r = p2.add_run(number)
        set_run_font(r, size=12)


def caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_together = True
    run = p.add_run(text)
    set_run_font(run, size=12, italic=True)


def note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    set_run_font(run, size=12, italic=True)


def add_picture(doc: Document, path: Path, title: str, width_cm: float = 15.0) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.first_line_indent = Cm(0)
    if path.is_file():
        p.add_run().add_picture(str(path), width=Cm(width_cm))
    else:
        run = p.add_run(f"[CHƯA CÓ HÌNH: {path.name}]")
        set_run_font(run, size=12, italic=True)
    caption(doc, title)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], title: str) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_run_font(run, size=11, bold=True)
        shade_cell(cell, "1F4E79")
        run.font.color.rgb = None
        # white text
        from docx.shared import RGBColor
        run.font.color.rgb = RGBColor(255, 255, 255)
    for r_i, row in enumerate(rows, start=1):
        for c_i, value in enumerate(row):
            cell = table.rows[r_i].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(value)
            set_run_font(run, size=11)
            if r_i % 2 == 0:
                shade_cell(cell, "F2F4F4")
    set_table_widths(table)
    caption(doc, title)


def shade_cell(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("w:shd"):
            tcPr.remove(child)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_table_widths(table) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), "5000")
    w.set(qn("w:type"), "pct")
    tblPr.append(w)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "autofit")
    tblPr.append(layout)


def page_break(doc: Document) -> None:
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)
